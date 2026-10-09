"""Excel template and transactional bulk product/size import."""
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from models import Product, Variant, TransactionType
from services import alert_service, inventory_service
from services.table_reader import read_table


MAX_IMPORT_ROWS = 20000
MAX_FILE_BYTES = 12 * 1024 * 1024

COLUMN_ALIASES = {
    "sku": {"sku", "style code", "sku / style code", "product sku"},
    "name": {"product name", "name", "design name", "design"},
    "category": {"category", "type"},
    "size": {"size", "size name", "variant size"},
    "stock": {"stock", "opening stock", "initial stock", "quantity", "qty"},
    "threshold": {"alert below", "alert threshold", "reorder threshold"},
    "target": {"target", "target stock", "target stock level"},
}


def _find_column(columns, aliases):
    normalized = {" ".join(str(column).strip().lower().split()): column for column in columns}
    for alias in aliases:
        if alias in normalized:
            return normalized[alias]
    return None


def _cell_text(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _nonnegative_int(value, *, default=None):
    text = _cell_text(value)
    if not text:
        if default is not None:
            return default
        raise ValueError("Stock is required.")
    try:
        number = Decimal(text)
    except (InvalidOperation, ValueError):
        raise ValueError(f"'{text}' must be a whole number 0 or greater.")
    if not number.is_finite() or number != number.to_integral_value() or number < 0:
        raise ValueError(f"'{text}' must be a whole number 0 or greater.")
    return int(number)


def build_bulk_product_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    headers = ["SKU", "Product Name", "Category", "Size", "Stock", "Alert Below", "Target Stock"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="18324F")
        cell.alignment = Alignment(horizontal="center")
    for row in range(2, 502):
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).number_format = "0"
    for column, width in {"A": 24, "B": 32, "C": 22, "D": 14, "E": 14, "F": 16, "G": 16}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{MAX_IMPORT_ROWS + 1}"

    help_ws = wb.create_sheet("Instructions")
    instructions = [
        ["Bulk product import"],
        ["Enter one row per size. Repeat SKU, Product Name and Category for every size of the same design."],
        ["Required columns: SKU, Product Name, Size and Stock. Stock must be a whole number 0 or greater."],
        ["Category may be blank. Alert Below defaults to 5 and Target Stock defaults to 20 when blank."],
        ["Do not include an image column. Add photos later from the product Edit screen."],
        ["A SKU already in Inventory or Trash is skipped. Fix the listed rows, then upload the corrected file."],
        ["If one row for a SKU is invalid, all sizes for that SKU are skipped together."],
        ["Duplicate SKU + Size rows and inconsistent names/categories are reported as row errors."],
        ["Supported files: .xlsx, .xls and .csv. Maximum 20,000 rows and 12 MB."],
    ]
    for row in instructions:
        help_ws.append(row)
    help_ws["A1"].font = Font(bold=True, color="FFFFFF", size=15)
    help_ws["A1"].fill = PatternFill("solid", fgColor="18324F")
    help_ws.column_dimensions["A"].width = 112
    for row in help_ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].font = Font(color="334155", size=11)
        help_ws.row_dimensions[row[0].row].height = 28

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_bulk_products(db: Session, file_bytes: bytes, filename: str = ""):
    if len(file_bytes) > MAX_FILE_BYTES:
        raise ValueError("File is larger than 12 MB. Split the inventory into smaller files.")
    try:
        df = read_table(file_bytes, filename)
    except ValueError:
        raise
    if df.empty:
        raise ValueError("The file has no product rows.")
    if len(df.index) > MAX_IMPORT_ROWS:
        raise ValueError(f"The file has {len(df.index)} rows. Upload at most {MAX_IMPORT_ROWS:,} at a time.")

    columns = {key: _find_column(df.columns, aliases) for key, aliases in COLUMN_ALIASES.items()}
    required = ("sku", "name", "size", "stock")
    missing = [key for key in required if columns[key] is None]
    if missing:
        names = {"sku": "SKU", "name": "Product Name", "size": "Size", "stock": "Stock"}
        raise ValueError("Missing required columns: " + ", ".join(names[key] for key in missing) + ". Download the template and copy your data into it.")

    groups: dict[str, list[dict]] = {}
    row_errors: dict[int, str] = {}
    row_skus: dict[int, str] = {}
    invalid_skus: set[str] = set()
    blank_rows = 0
    for index, record in enumerate(df.to_dict("records"), start=2):
        values = [_cell_text(value) for value in record.values()]
        if not any(values):
            blank_rows += 1
            continue

        sku = _cell_text(record.get(columns["sku"])).upper()
        if not sku:
            row_errors[index] = "SKU is required."
            continue
        row_skus[index] = sku
        groups.setdefault(sku, [])
        try:
            name = _cell_text(record.get(columns["name"]))
            size = _cell_text(record.get(columns["size"])).upper()
            if not name:
                raise ValueError("Product Name is required.")
            if not size:
                raise ValueError("Size is required.")
            if len(sku) > 64 or len(name) > 300 or len(size) > 20:
                raise ValueError("SKU, Product Name or Size exceeds the allowed length.")
            category = _cell_text(record.get(columns["category"])) if columns["category"] else ""
            if len(category) > 100:
                raise ValueError("Category must be 100 characters or fewer.")
            stock = _nonnegative_int(record.get(columns["stock"]))
            threshold = _nonnegative_int(record.get(columns["threshold"]), default=5) if columns["threshold"] else 5
            target = _nonnegative_int(record.get(columns["target"]), default=20) if columns["target"] else 20
            groups[sku].append({
                "row": index, "sku": sku, "name": name, "category": category,
                "size": size, "stock": stock, "threshold": threshold, "target": target,
            })
        except ValueError as exc:
            row_errors[index] = str(exc)
            invalid_skus.add(sku)

    # Validate each SKU as one unit so a typo on one size never leaves a
    # partially-created design in the catalog.
    group_errors: dict[str, str] = {}
    for sku, entries in groups.items():
        if sku in invalid_skus:
            group_errors[sku] = "One or more rows for this SKU are invalid; all sizes for this SKU were skipped."
            continue
        names = {entry["name"].casefold() for entry in entries}
        categories = {entry["category"].casefold() for entry in entries if entry["category"]}
        sizes = [entry["size"] for entry in entries]
        if len(names) > 1:
            group_errors[sku] = "Product Name must match on every row for this SKU."
        elif len(categories) > 1:
            group_errors[sku] = "Category must match on every row for this SKU."
        elif len(sizes) != len(set(sizes)):
            group_errors[sku] = "This SKU has a repeated size. Keep only one row per SKU + Size."

    valid_skus = [sku for sku in groups if sku not in group_errors]
    existing_products = {
        product.sku: product for product in db.query(Product).filter(Product.sku.in_(valid_skus)).all()
    } if valid_skus else {}
    for sku, product in existing_products.items():
        if product.is_deleted:
            group_errors[sku] = "This SKU is in Trash. Restore it before importing this SKU."
        else:
            group_errors[sku] = "This SKU already exists. Use Edit to add sizes or update its stock."

    valid_groups = [(sku, entries) for sku, entries in groups.items() if sku not in group_errors]
    created_products = created_variants = 0
    now = datetime.now(timezone.utc)
    try:
        for sku, entries in valid_groups:
            first = entries[0]
            category = next((entry["category"] for entry in entries if entry["category"]), "")
            product = Product(sku=sku, name=first["name"], category=category or None)
            db.add(product)
            db.flush()
            created_products += 1
            for entry in entries:
                variant = Variant(
                    product_id=product.id,
                    size=entry["size"],
                    variant_code=f"{sku}_{entry['size']}",
                    current_stock=0,
                    reorder_threshold=entry["threshold"],
                    target_stock_level=entry["target"],
                    last_counted_at=now,
                )
                db.add(variant)
                db.flush()
                if entry["stock"]:
                    inventory_service.apply_stock_change(
                        db, variant, entry["stock"],
                        transaction_type=TransactionType.INITIAL_STOCK,
                        reference="bulk product import", created_by="bulk_product_import",
                    )
                alert_service.check_variant_alert(db, variant)
                created_variants += 1
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise ValueError("Import could not be saved because a SKU or size already exists. Refresh the inventory and upload again.") from exc

    errors = []
    for row, message in row_errors.items():
        errors.append({"row": row, "sku": row_skus.get(row), "message": message})
    for sku, message in group_errors.items():
        for entry in groups[sku]:
            if entry["row"] not in row_errors:
                errors.append({"row": entry["row"], "sku": sku, "message": message})
    errors.sort(key=lambda item: item["row"])

    processed_rows = len(df.index)
    rows_skipped = processed_rows - blank_rows - created_variants
    return {
        "rows_processed": processed_rows,
        "products_created": created_products,
        "variants_created": created_variants,
        "rows_skipped": rows_skipped,
        "blank_rows": blank_rows,
        "errors": errors[:100],
        "errors_truncated": len(errors) > 100,
    }
