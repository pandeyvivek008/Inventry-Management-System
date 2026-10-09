"""Marketplace-order shortage picklist.

Pending orders reserve current stock until dispatch. The picklist reports
pending demand, current physical stock, dispatched units for context, and the
remaining quantity to restock. Dispatched units are not subtracted twice.
"""
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, case
from sqlalchemy.orm import Session, selectinload

from models import Variant, Product, ReorderBatch, ReorderBatchItem, OrderLine, OrderLineStatus


def generate_reorder_list(db: Session, save_batch: bool = True, part: str | None = None):
    """Return one row per stockable size with replenishment demand.

    Pending orders reserve physical stock until dispatch. Dispatched counts are
    shown separately for context, since those quantities have already reduced
    physical stock. Required is based only on pending orders.
    """
    query = (
        db.query(
            Variant,
            func.coalesce(func.sum(case(
                (OrderLine.status == OrderLineStatus.PENDING, OrderLine.qty_ordered), else_=0
            )), 0).label("ordered_qty"),
            func.coalesce(func.sum(case(
                (OrderLine.status == OrderLineStatus.FULFILLED, OrderLine.qty_ordered), else_=0
            )), 0).label("dispatched_qty"),
        )
        .join(Product, Variant.product_id == Product.id)
        .options(selectinload(Variant.product))
        .outerjoin(
            OrderLine,
            (OrderLine.variant_id == Variant.id)
            & (OrderLine.status.in_([OrderLineStatus.PENDING, OrderLineStatus.FULFILLED])),
        )
        .filter(
            Variant.last_counted_at.isnot(None),
            Variant.is_deleted == False,  # noqa: E712
            Product.is_deleted == False,  # noqa: E712
        )
        .group_by(Variant.id)
    )
    if part and part.strip():
        query = query.filter(Variant.variant_code.ilike(f"%{part.strip().upper()}%"))

    rows = query.all()
    items = []
    selected = []
    for v, ordered_qty, dispatched_qty in rows:
        ordered_qty = int(ordered_qty or 0)
        dispatched_qty = int(dispatched_qty or 0)
        # Only pending orders reserve current stock. Dispatched orders have
        # already reduced physical stock and must not create a second shortage.
        if ordered_qty <= 0 and dispatched_qty <= 0:
            continue
        physical_stock = int(v.current_stock or 0)
        available_stock = max(physical_stock, 0)
        required = max(ordered_qty - available_stock, 0)
        items.append({
            "sku": v.product.sku,
            "size": v.size.upper(),
            "variant_code": v.variant_code,
            "name": v.product.name,
            "current_stock": physical_stock,
            "ordered_qty": ordered_qty,
            "dispatched_qty": dispatched_qty,
            "required_qty": required,
            "image_path": v.product.image_path,
            "is_new": False,
        })
        if required > 0:
            selected.append((v, required))

    # Orders that could not be mapped to an existing product must NOT create
    # a product automatically. They are surfaced as NEW in the picklist so
    # the team can identify the SKU and arrange the quantity manually.
    unknown_query = (
        db.query(
            OrderLine.sku, OrderLine.size,
            func.sum(OrderLine.qty_ordered).label("ordered_qty"),
        )
        .filter(OrderLine.status == OrderLineStatus.SKU_NOT_FOUND)
        .group_by(OrderLine.sku, OrderLine.size)
    )
    if part and part.strip():
        unknown_query = unknown_query.filter(OrderLine.sku.ilike(f"%{part.strip().upper()}%"))
    for sku, size, ordered_qty in unknown_query.all():
        ordered_qty = int(ordered_qty or 0)
        if ordered_qty <= 0:
            continue
        items.append({
            "sku": str(sku).upper(),
            "size": str(size or "—").upper(),
            "variant_code": str(sku).upper(),
            "name": "NEW / UNMAPPED SKU",
            "current_stock": 0,
            "ordered_qty": ordered_qty,
            "dispatched_qty": 0,
            "required_qty": ordered_qty,
            "image_path": None,
            "is_new": True,
        })

    items.sort(key=lambda x: (-x["required_qty"], x["current_stock"], x["sku"]))

    if save_batch and selected:
        batch = ReorderBatch()
        db.add(batch)
        db.flush()
        for v, required in selected:
            db.add(ReorderBatchItem(
                batch_id=batch.id,
                variant_id=v.id,
                stock_at_generation=v.current_stock,
                required_qty=required,
            ))
        db.commit()

    return items


def reorder_list_to_excel(items: list[dict]) -> bytes:
    """Purchase/picklist export: ONLY SKU (size included) + Quantity."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Picklist"

    headers = ["SKU", "Size", "Quantity"]
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F3B52", end_color="2F3B52", fill_type="solid")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, item in enumerate(items, start=2):
        ws.cell(row=row_idx, column=1, value=item["sku"]).font = Font(name="Arial")
        ws.cell(row=row_idx, column=2, value=item["size"]).font = Font(name="Arial")
        ws.cell(row=row_idx, column=3, value=int(item["required_qty"])).font = Font(name="Arial")

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{max(2, len(items) + 1)}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
